"""
Écran de saisie type DHIS2 (Tranches 2-3).

L'utilisateur choisit éventuellement un formulaire (data_set), puis une unité
d'organisation et une période. La grille affiche une ligne par élément de
données × désagrégation :
  - avec formulaire : seuls les éléments du formulaire, périodes de son type ;
  - sans formulaire (saisie libre) : tous les éléments, périodes mensuelles.
À l'enregistrement, chaque cellule non vide est écrite dans DataValue
(upsert idempotent) ; une cellule vidée supprime la valeur existante.
"""
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.shortcuts import render, redirect

from metadata.models import OrgUnit, DataElement, CategoryOptionCombo
from .models import DataValue, DataSet
from . import periods


def _leaf_org_units():
    """Unités de saisie : les feuilles de la hiérarchie (formations sanitaires)."""
    return OrgUnit.objects.filter(children__isnull=True).order_by("name")


def _elements_for(dataset):
    """Éléments à afficher : ceux du formulaire, ou tous si saisie libre."""
    if dataset:
        return dataset.data_elements.select_related("category_combo").order_by("name")
    return DataElement.objects.select_related("category_combo").order_by("name")


def _org_units_for(dataset):
    """Unités proposées : celles du formulaire (si définies), sinon les feuilles."""
    if dataset and dataset.org_units.exists():
        return dataset.org_units.order_by("name")
    return _leaf_org_units()


def _build_rows(elements, org_unit, period):
    """Lignes de la grille : par élément, ses désagrégations, valeur pré-remplie."""
    existing = {
        (v.data_element_id, v.category_option_combo_id): v.value
        for v in DataValue.objects.filter(org_unit=org_unit, period=period)
    }
    rows = []
    for de in elements:
        cocs = CategoryOptionCombo.objects.filter(combo=de.category_combo).order_by("name")
        cells = []
        for coc in cocs:
            cells.append({
                "field_name": f"de{de.id}_coc{coc.id}",
                "coc": coc,
                "value": existing.get((de.id, coc.id), ""),
            })
        rows.append({"element": de, "cells": cells})
    return rows


@login_required
def data_entry(request):
    datasets = DataSet.objects.order_by("name")

    src = request.POST if request.method == "POST" else request.GET
    ds_id = src.get("dataset") or ""
    ou_id = src.get("org_unit")
    period = src.get("period")

    dataset = None
    if ds_id:
        dataset = DataSet.objects.filter(pk=ds_id).first()

    ptype = dataset.period_type if dataset else periods.MONTHLY
    period_choices = [(c, periods.label(c)) for c in periods.recent_codes(ptype, count=12)]

    context = {
        "datasets": datasets,
        "selected_dataset": dataset,
        "org_units": _org_units_for(dataset),
        "period_choices": period_choices,
        "selected_ou": None,
        "selected_period": period,
        "rows": None,
    }

    if not ou_id or not period:
        return render(request, "collecte/data_entry.html", context)

    try:
        org_unit = OrgUnit.objects.get(pk=ou_id)
        periods.validate(period, expected_type=ptype if dataset else None)
    except (OrgUnit.DoesNotExist, periods.InvalidPeriod):
        messages.error(request, "Unité d'organisation ou période invalide pour ce formulaire.")
        return render(request, "collecte/data_entry.html", context)

    context["selected_ou"] = org_unit
    elements = _elements_for(dataset)

    if request.method == "POST" and "save" in request.POST:
        saved, cleared, errors = _save_grid(request, elements, org_unit, period)
        for err in errors:
            messages.error(request, err)
        if saved or cleared:
            messages.success(
                request,
                f"{saved} valeur(s) enregistrée(s), {cleared} supprimée(s) "
                f"pour {org_unit.name} — {periods.label(period)}."
            )
        qs = f"?org_unit={org_unit.id}&period={period}"
        if dataset:
            qs += f"&dataset={dataset.id}"
        return redirect(f"{request.path}{qs}")

    context["rows"] = _build_rows(elements, org_unit, period)
    return render(request, "collecte/data_entry.html", context)


@transaction.atomic
def _save_grid(request, elements, org_unit, period):
    """Traite les cellules soumises. Retourne (nb_enregistrées, nb_supprimées, erreurs)."""
    saved = cleared = 0
    errors = []
    for de in elements:
        cocs = CategoryOptionCombo.objects.filter(combo=de.category_combo)
        for coc in cocs:
            field_name = f"de{de.id}_coc{coc.id}"
            if field_name not in request.POST:
                continue
            raw = request.POST.get(field_name, "").strip()

            if raw == "":
                deleted, _ = DataValue.objects.filter(
                    data_element=de, org_unit=org_unit,
                    period=period, category_option_combo=coc,
                ).delete()
                cleared += deleted
                continue

            try:
                value = Decimal(raw.replace(",", "."))
            except (InvalidOperation, ValueError):
                errors.append(f"« {de.name} » : valeur non numérique ({raw!r}).")
                continue

            dv = DataValue(
                data_element=de, org_unit=org_unit, period=period,
                category_option_combo=coc, value=value,
            )
            try:
                # Validation métier SANS la contrainte d'unicité : l'upsert
                # ci-dessous gère le cas « la valeur existe déjà » (mise à jour).
                dv.clean()
            except ValidationError as e:
                msg = "; ".join(
                    m for msgs in getattr(e, "message_dict", {}).values() for m in msgs
                ) or str(e)
                errors.append(f"« {de.name} » : {msg}")
                continue

            DataValue.objects.update_or_create(
                data_element=de, org_unit=org_unit,
                period=period, category_option_combo=coc,
                defaults={"value": value},
            )
            saved += 1
    return saved, cleared, errors


@login_required
def download_template(request):
    """Génère et renvoie le template .xlsx pour un formulaire (ou saisie libre)."""
    from django.http import HttpResponse
    from .excelio import build_template
    from .models import DataSet

    ds_id = request.GET.get("dataset") or ""
    period = request.GET.get("period") or None
    dataset = DataSet.objects.filter(pk=ds_id).first() if ds_id else None

    data = build_template(dataset, period=period)
    fname = f"template_{dataset.code if dataset else 'libre'}"
    if period:
        fname += f"_{period}"
    fname += ".xlsx"

    resp = HttpResponse(
        data,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="{fname}"'
    return resp


@login_required
def import_excel(request):
    """Page d'import : upload d'un .xlsx rempli, affichage du rapport d'anomalies."""
    from .excelio import parse_and_import
    from .models import DataSet

    datasets = DataSet.objects.order_by("name")
    context = {"datasets": datasets, "report": None,
               "selected_dataset": None, "selected_period": None}

    if request.method == "POST":
        ds_id = request.POST.get("dataset") or ""
        period = request.POST.get("period") or ""
        dataset = DataSet.objects.filter(pk=ds_id).first() if ds_id else None
        dry = request.POST.get("dry_run") == "1"
        upload = request.FILES.get("file")

        context["selected_dataset"] = dataset
        context["selected_period"] = period

        if not upload:
            messages.error(request, "Veuillez sélectionner un fichier .xlsx.")
        elif not period:
            messages.error(request, "Veuillez indiquer la période.")
        else:
            report = parse_and_import(upload, dataset, period, save=not dry)
            context["report"] = report.as_dict()
            if dry:
                messages.info(request, f"Simulation : {report.saved} valeur(s) seraient enregistrées.")
            elif report.ok:
                messages.success(request, f"{report.saved} valeur(s) importée(s), {report.cleared} supprimée(s).")
            else:
                messages.warning(request, f"Import terminé avec anomalies : {report.saved} enregistrée(s), voir le détail.")

    # Périodes proposées (12 derniers mois par défaut)
    context["period_choices"] = [(c, periods.label(c)) for c in periods.recent_codes(count=12)]
    return render(request, "collecte/import_excel.html", context)


@login_required
def import_flexible(request):
    """
    Import souple (mode fichier libre) en deux étapes :
      1. upload + analyse : on montre les colonnes inconnues et on demande une
         décision pour chacune (rattacher à un élément existant / créer / ignorer) ;
      2. confirmation : on applique les décisions puis on importe.
    Le fichier est conservé temporairement entre les deux étapes.
    """
    import os
    import tempfile
    import uuid

    from django.conf import settings
    from .excelio import analyze_columns, create_element_from_column, parse_and_import
    from .models import DataSet
    from metadata.models import DataElement

    tmp_dir = os.path.join(tempfile.gettempdir(), "sismalaria_imports")
    os.makedirs(tmp_dir, exist_ok=True)

    datasets = DataSet.objects.order_by("name")
    all_elements = DataElement.objects.order_by("name")
    context = {
        "datasets": datasets,
        "all_elements": all_elements,
        "period_choices": [(c, periods.label(c)) for c in periods.recent_codes(count=12)],
        "step": "upload",
    }

    # --- Étape 2 : confirmation (les décisions arrivent) ---
    if request.method == "POST" and request.POST.get("confirm") == "1":
        token = request.POST.get("token", "")
        period = request.POST.get("period", "")
        ds_id = request.POST.get("dataset") or ""
        dataset = DataSet.objects.filter(pk=ds_id).first() if ds_id else None
        path = os.path.join(tmp_dir, f"{token}.xlsx")

        if not token or not os.path.exists(path):
            messages.error(request, "Session d'import expirée, veuillez recommencer.")
            return render(request, "collecte/import_flexible.html", context)

        # Appliquer les décisions sur les colonnes inconnues
        created = []
        analysis = analyze_columns(open(path, "rb"), dataset)

        # Décisions par index de colonne : on crée les éléments demandés et on
        # note la nouvelle clé technique à réécrire dans le fichier.
        key_rewrites = {}  # {col_index: "CODE::coc_id"}
        for col in analysis["unknown"]:
            idx = col["index"]
            action = request.POST.get(f"action_{idx}", "ignore")
            if action == "create":
                name = request.POST.get(f"name_{idx}", "").strip() or col["header"]
                code = request.POST.get(f"code_{idx}", "").strip()
                if not code:
                    messages.error(request, f"Code manquant pour « {name} ».")
                    continue
                de = create_element_from_column(name, code)
                created.append(de.code)
                if dataset:
                    dataset.data_elements.add(de)
                # Récupère le coc par défaut du nouvel élément pour la clé
                from metadata.models import CategoryOptionCombo as COC
                coc = COC.objects.filter(combo=de.category_combo).first()
                if coc:
                    from .excelio import KEY_SEP
                    key_rewrites[idx] = f"{de.code}{KEY_SEP}{coc.id}"

        # Réécrit les clés des colonnes créées dans le fichier temporaire,
        # afin que l'import les reconnaisse (le placeholder ::0 devient ::<coc_id>).
        if key_rewrites:
            from openpyxl import load_workbook
            from .excelio import KEY_ROW
            wbx = load_workbook(path)
            wsx = wbx["Données"]
            for idx, new_key in key_rewrites.items():
                wsx.cell(KEY_ROW, idx + 1, new_key)  # openpyxl est 1-based
            wbx.save(path)

        # Import : en mode souple, on importe sans formulaire pour accepter
        # toutes les colonnes désormais connues (les nouveaux éléments inclus)
        report = parse_and_import(open(path, "rb"), None, period, save=True)
        rep = report.as_dict()
        rep["created_elements"] = created
        context["report"] = rep
        context["step"] = "done"

        # Nettoyage du fichier temporaire
        try:
            os.remove(path)
        except OSError:
            pass

        if created:
            messages.success(request, f"{len(created)} nouvel(s) élément(s) créé(s) : {', '.join(created)}.")
        if report.ok or report.saved:
            messages.success(request, f"{report.saved} valeur(s) importée(s).")
        else:
            messages.warning(request, "Import terminé, voir le rapport.")
        return render(request, "collecte/import_flexible.html", context)

    # --- Étape 1 : upload + analyse ---
    if request.method == "POST" and request.POST.get("analyze") == "1":
        upload = request.FILES.get("file")
        period = request.POST.get("period", "")
        ds_id = request.POST.get("dataset") or ""
        dataset = DataSet.objects.filter(pk=ds_id).first() if ds_id else None

        if not upload:
            messages.error(request, "Veuillez sélectionner un fichier .xlsx.")
            return render(request, "collecte/import_flexible.html", context)
        if not period:
            messages.error(request, "Veuillez indiquer la période.")
            return render(request, "collecte/import_flexible.html", context)

        # Sauvegarde temporaire
        token = uuid.uuid4().hex
        path = os.path.join(tmp_dir, f"{token}.xlsx")
        with open(path, "wb") as f:
            for chunk in upload.chunks():
                f.write(chunk)

        analysis = analyze_columns(open(path, "rb"), dataset)
        if analysis["error"]:
            messages.error(request, analysis["error"])
            try:
                os.remove(path)
            except OSError:
                pass
            return render(request, "collecte/import_flexible.html", context)

        context.update({
            "step": "review",
            "token": token,
            "selected_period": period,
            "selected_dataset": dataset,
            "analysis": analysis,
            "n_known": len(analysis["known"]),
            "n_unknown": len(analysis["unknown"]),
        })
        return render(request, "collecte/import_flexible.html", context)

    return render(request, "collecte/import_flexible.html", context)


@login_required
def import_dhis2_page(request):
    """Page dédiée à l'import d'un export DHIS2 (gros fichier, format large)."""
    import os, tempfile
    from .dhis2_import import import_dhis2

    context = {"report": None}

    if request.method == "POST":
        upload = request.FILES.get("file")
        sheet = request.POST.get("sheet_name", "").strip() or 0
        dry = request.POST.get("dry_run") == "1"
        create_ou = request.POST.get("create_ou") == "1"

        if not upload:
            messages.error(request, "Veuillez sélectionner un fichier .xlsx.")
            return render(request, "collecte/import_dhis2.html", context)

        # Sauvegarde temporaire (gros fichier)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        for chunk in upload.chunks():
            tmp.write(chunk)
        tmp.close()

        try:
            report = import_dhis2(tmp.name, sheet_name=sheet, save=not dry,
                                  create_missing_org_units=create_ou)
            context["report"] = report.as_dict()
            if dry:
                messages.info(request,
                    f"Simulation : {context['report']['columns_ok']} colonnes reconnues, "
                    f"{context['report']['value_cells']} valeurs seraient importées.")
            elif report.errors:
                messages.warning(request, "Import terminé avec des erreurs, voir le rapport.")
            else:
                messages.success(request, f"{report.saved} valeurs importées.")
        except Exception as e:
            messages.error(request, f"Erreur pendant l'import : {e}")
        finally:
            os.unlink(tmp.name)

    return render(request, "collecte/import_dhis2.html", context)


# ---------------------------------------------------------------------------
#  Accueil des imports (brique 5a) — liste les types d'import disponibles
# ---------------------------------------------------------------------------
def import_home(request):
    """Page d'accueil de la section Imports : une tuile par type d'import."""
    types_import = [
        {
            "titre": "Extraction DHIS2 (routine)",
            "desc": "Importer un export DHIS2 en largeur (données de routine paludisme).",
            "url_name": "collecte:import_dhis2",
            "dispo": True,
        },
        {
            "titre": "Population (RGPH)",
            "desc": "Charger la population de référence par district et par année (dénominateur).",
            "url_name": "collecte:import_population",
            "dispo": True,
        },
        {
            "titre": "Cartes (shapefiles)",
            "desc": "Charger les limites géographiques pour la stratification (brique cartographie).",
            "url_name": None,
            "dispo": False,
        },
        {
            "titre": "Données épidémiologiques",
            "desc": "Import d'autres fichiers Excel (format à définir avec la hiérarchie).",
            "url_name": None,
            "dispo": False,
        },
        {
            "titre": "Import par template (strict)",
            "desc": "Import Excel encadré par un modèle généré.",
            "url_name": "collecte:import_excel",
            "dispo": True,
        },
        {
            "titre": "Import souple (colonnes libres)",
            "desc": "Import Excel à colonnes libres avec validation humaine.",
            "url_name": "collecte:import_flexible",
            "dispo": True,
        },
    ]
    return render(request, "collecte/import_home.html", {
        "active": "imports",
        "types_import": types_import,
    })


# ---------------------------------------------------------------------------
#  Import Population (brique 5b) — page web branchée sur la brique 1
# ---------------------------------------------------------------------------
def import_population_page(request):
    """Page dédiée à l'import du fichier population (RGPH → table Population)."""
    import os, tempfile
    from .population_service import import_population

    context = {"report": None}

    if request.method == "POST":
        upload = request.FILES.get("file")
        dry = request.POST.get("dry_run") == "1"

        if not upload:
            messages.error(request, "Veuillez sélectionner un fichier .xlsx.")
            return render(request, "collecte/import_population.html", context)

        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        for chunk in upload.chunks():
            tmp.write(chunk)
        tmp.close()

        try:
            report = import_population(tmp.name, save=not dry)
            context["report"] = report.as_dict()
            if dry:
                messages.info(request,
                    f"Simulation : {report.matched_districts} districts reconnus "
                    f"sur {report.parsed_rows} lignes lues.")
            elif report.unmatched_districts:
                messages.warning(request,
                    f"{report.saved} lignes importées. "
                    f"{len(report.unmatched_districts)} district(s) non reconnu(s) — voir le rapport.")
            else:
                messages.success(request,
                    f"{report.saved} lignes importées, {report.matched_districts} districts.")
        except Exception as e:
            messages.error(request, f"Erreur pendant l'import : {e}")
        finally:
            os.unlink(tmp.name)

    return render(request, "collecte/import_population.html", context)
