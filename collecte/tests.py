"""Tests de la tranche 2 (collecte / valeurs)."""
from decimal import Decimal

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.test import TestCase, Client
from django.urls import reverse

from metadata.models import (
    OrgUnit, Category, CategoryOption, CategoryCombo,
    CategoryOptionCombo, DataElement,
)
from .models import DataValue
from . import periods


class PeriodTests(TestCase):
    def test_detect_type(self):
        self.assertEqual(periods.detect_type("202608"), periods.MONTHLY)
        self.assertEqual(periods.detect_type("2026Q3"), periods.QUARTERLY)
        self.assertEqual(periods.detect_type("2026"), periods.YEARLY)

    def test_invalid_period(self):
        with self.assertRaises(periods.InvalidPeriod):
            periods.detect_type("2026-08")
        with self.assertRaises(periods.InvalidPeriod):
            periods.detect_type("202613")  # mois 13

    def test_label(self):
        self.assertEqual(periods.label("202608"), "août 2026")
        self.assertEqual(periods.label("2026Q3"), "T3 2026")

    def test_recent_codes(self):
        import datetime
        codes = periods.recent_codes(count=3, ref=datetime.date(2026, 8, 15))
        self.assertEqual(codes, ["202608", "202607", "202606"])


class DataValueModelTests(TestCase):
    def setUp(self):
        self.pays = OrgUnit.objects.create(name="Togo", code="TG", level=1)
        self.fs = OrgUnit.objects.create(name="USP Agoè", code="AGOE", parent=self.pays)
        self.combo = CategoryCombo.objects.create(name="(défaut)", code="default", is_default=True)
        self.coc = CategoryOptionCombo.objects.create(combo=self.combo, name="(défaut)")
        self.de = DataElement.objects.create(
            name="Cas confirmés", code="PAL_CONF", category_combo=self.combo,
        )

    def test_create_value(self):
        dv = DataValue.objects.create(
            data_element=self.de, org_unit=self.fs, period="202608",
            category_option_combo=self.coc, value=Decimal("342"),
        )
        self.assertIn("PAL_CONF@AGOE/202608 = 342", str(dv))
        self.assertEqual(dv.period_label, "août 2026")

    def test_uniqueness_constraint(self):
        DataValue.objects.create(
            data_element=self.de, org_unit=self.fs, period="202608",
            category_option_combo=self.coc, value=Decimal("342"),
        )
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                DataValue.objects.create(
                    data_element=self.de, org_unit=self.fs, period="202608",
                    category_option_combo=self.coc, value=Decimal("999"),
                )

    def test_invalid_period_rejected(self):
        dv = DataValue(
            data_element=self.de, org_unit=self.fs, period="2026-08",
            category_option_combo=self.coc, value=Decimal("10"),
        )
        with self.assertRaises(ValidationError):
            dv.full_clean()

    def test_wrong_disaggregation_rejected(self):
        # Un coc d'un autre combo ne doit pas être accepté
        other = CategoryCombo.objects.create(name="Autre", code="autre")
        other_coc = CategoryOptionCombo.objects.create(combo=other, name="X")
        dv = DataValue(
            data_element=self.de, org_unit=self.fs, period="202608",
            category_option_combo=other_coc, value=Decimal("10"),
        )
        with self.assertRaises(ValidationError):
            dv.full_clean()

    def test_negative_rejected_for_positive_integer(self):
        dv = DataValue(
            data_element=self.de, org_unit=self.fs, period="202608",
            category_option_combo=self.coc, value=Decimal("-5"),
        )
        with self.assertRaises(ValidationError):
            dv.full_clean()


class DataEntryViewTests(TestCase):
    def setUp(self):
        self.pays = OrgUnit.objects.create(name="Togo", code="TG", level=1)
        self.fs = OrgUnit.objects.create(name="USP Agoè", code="AGOE", parent=self.pays)
        self.combo = CategoryCombo.objects.create(name="(défaut)", code="default", is_default=True)
        self.coc = CategoryOptionCombo.objects.create(combo=self.combo, name="(défaut)")
        self.de = DataElement.objects.create(
            name="Cas confirmés", code="PAL_CONF", category_combo=self.combo,
        )
        self.user = User.objects.create_superuser("admin", "a@b.c", "pwd12345")
        self.client = Client()
        self.client.force_login(self.user)

    def test_grid_displays(self):
        url = reverse("collecte:data_entry")
        resp = self.client.get(url, {"org_unit": self.fs.id, "period": "202608"})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Cas confirmés")

    def test_save_creates_value(self):
        url = reverse("collecte:data_entry")
        field = f"de{self.de.id}_coc{self.coc.id}"
        resp = self.client.post(url, {
            "org_unit": self.fs.id, "period": "202608",
            "save": "1", field: "342",
        })
        self.assertEqual(resp.status_code, 302)  # redirect après save
        dv = DataValue.objects.get(data_element=self.de, org_unit=self.fs, period="202608")
        self.assertEqual(dv.value, Decimal("342"))

    def test_save_is_idempotent_update(self):
        url = reverse("collecte:data_entry")
        field = f"de{self.de.id}_coc{self.coc.id}"
        self.client.post(url, {"org_unit": self.fs.id, "period": "202608", "save": "1", field: "342"})
        self.client.post(url, {"org_unit": self.fs.id, "period": "202608", "save": "1", field: "400"})
        self.assertEqual(DataValue.objects.filter(data_element=self.de, org_unit=self.fs, period="202608").count(), 1)
        dv = DataValue.objects.get(data_element=self.de, org_unit=self.fs, period="202608")
        self.assertEqual(dv.value, Decimal("400"))

    def test_empty_clears_value(self):
        url = reverse("collecte:data_entry")
        field = f"de{self.de.id}_coc{self.coc.id}"
        self.client.post(url, {"org_unit": self.fs.id, "period": "202608", "save": "1", field: "342"})
        self.client.post(url, {"org_unit": self.fs.id, "period": "202608", "save": "1", field: ""})
        self.assertEqual(DataValue.objects.filter(data_element=self.de, org_unit=self.fs, period="202608").count(), 0)

    def test_login_required(self):
        self.client.logout()
        url = reverse("collecte:data_entry")
        resp = self.client.get(url, {"org_unit": self.fs.id, "period": "202608"})
        self.assertEqual(resp.status_code, 302)  # redirige vers login


class DataSetTests(TestCase):
    def setUp(self):
        self.pays = OrgUnit.objects.create(name="Togo", code="TG", level=1)
        self.fs = OrgUnit.objects.create(name="USP Agoè", code="AGOE", parent=self.pays)
        self.combo = CategoryCombo.objects.create(name="(défaut)", code="default", is_default=True)
        self.coc = CategoryOptionCombo.objects.create(combo=self.combo, name="(défaut)")
        self.de1 = DataElement.objects.create(name="Cas confirmés", code="C1", category_combo=self.combo)
        self.de2 = DataElement.objects.create(name="Cas traités", code="C2", category_combo=self.combo)
        self.de3 = DataElement.objects.create(name="MILDA", code="C3", category_combo=self.combo)
        from collecte.models import DataSet
        self.ds = DataSet.objects.create(name="Rapport mensuel", code="RM", period_type=periods.MONTHLY)
        self.ds.data_elements.set([self.de1, self.de2])
        self.ds.org_units.set([self.fs])
        self.user = User.objects.create_superuser("admin2", "a@b.c", "pwd12345")
        self.client = Client()
        self.client.force_login(self.user)

    def test_dataset_holds_elements_and_units(self):
        self.assertEqual(self.ds.data_elements.count(), 2)
        self.assertEqual(self.ds.org_units.count(), 1)

    def test_grid_with_dataset_shows_only_its_elements(self):
        url = reverse("collecte:data_entry")
        resp = self.client.get(url, {"dataset": self.ds.id, "org_unit": self.fs.id, "period": "202608"})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Cas confirmés")
        self.assertContains(resp, "Cas traités")
        self.assertNotContains(resp, "MILDA")  # hors formulaire

    def test_grid_without_dataset_shows_all(self):
        url = reverse("collecte:data_entry")
        resp = self.client.get(url, {"org_unit": self.fs.id, "period": "202608"})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "MILDA")  # saisie libre : tout

    def test_period_type_mismatch_rejected(self):
        # Le formulaire est mensuel : une période trimestrielle doit être refusée
        url = reverse("collecte:data_entry")
        resp = self.client.get(url, {"dataset": self.ds.id, "org_unit": self.fs.id, "period": "2026Q3"})
        # La grille ne se construit pas (période invalide pour ce type)
        self.assertIsNone(resp.context["rows"])

    def test_save_via_dataset(self):
        url = reverse("collecte:data_entry")
        field = f"de{self.de1.id}_coc{self.coc.id}"
        resp = self.client.post(url, {
            "dataset": self.ds.id, "org_unit": self.fs.id, "period": "202608",
            "save": "1", field: "150",
        })
        self.assertEqual(resp.status_code, 302)
        from collecte.models import DataValue
        dv = DataValue.objects.get(data_element=self.de1, org_unit=self.fs, period="202608")
        self.assertEqual(dv.value, Decimal("150"))


class ExcelImportTests(TestCase):
    def setUp(self):
        from collecte.models import DataSet
        self.pays = OrgUnit.objects.create(name="Togo", code="TG", level=1)
        self.fs1 = OrgUnit.objects.create(name="CHU", code="CHU", parent=self.pays)
        self.fs2 = OrgUnit.objects.create(name="USP Agoè", code="AGOE", parent=self.pays)
        self.combo = CategoryCombo.objects.create(name="(défaut)", code="default", is_default=True)
        self.coc = CategoryOptionCombo.objects.create(combo=self.combo, name="(défaut)")
        self.de1 = DataElement.objects.create(name="Cas suspects", code="SUSP", category_combo=self.combo)
        self.de2 = DataElement.objects.create(name="TDR positifs", code="TDRP", category_combo=self.combo)
        self.ds = DataSet.objects.create(name="Rapport mensuel", code="RM", period_type=periods.MONTHLY)
        self.ds.data_elements.set([self.de1, self.de2])
        self.ds.org_units.set([self.fs1, self.fs2])
        self.user = User.objects.create_superuser("adm3", "a@b.c", "pwd12345")
        self.client = Client()
        self.client.force_login(self.user)

    def _fill_and_get_buffer(self, values):
        """Génère le template, le remplit selon values={(ou_code,de_code):val}, renvoie un buffer."""
        from collecte.excelio import build_template, KEY_SEP
        from openpyxl import load_workbook
        from io import BytesIO
        data = build_template(self.ds, period="202608")
        wb = load_workbook(BytesIO(data))
        ws = wb["Données"]
        keys = {ws.cell(2, c).value: c for c in range(1, ws.max_column + 1)}
        ou_rows = {ws.cell(r, 1).value: r for r in range(3, ws.max_row + 1)}
        for (ou_code, de_code), val in values.items():
            # trouver la colonne dont la clé commence par de_code::
            col = next(c for k, c in keys.items() if k and k.startswith(f"{de_code}{KEY_SEP}"))
            ws.cell(ou_rows[ou_code], col, val)
        buf = BytesIO(); wb.save(buf); buf.seek(0)
        return buf

    def test_build_template_structure(self):
        from collecte.excelio import build_template
        from openpyxl import load_workbook
        from io import BytesIO
        data = build_template(self.ds, period="202608")
        wb = load_workbook(BytesIO(data))
        self.assertIn("Données", wb.sheetnames)
        self.assertIn("Informations", wb.sheetnames)
        ws = wb["Données"]
        # 2 colonnes fixes + 2 éléments (combo défaut = 1 coc chacun) = 4 colonnes
        self.assertEqual(ws.max_column, 4)
        # 2 unités
        self.assertEqual(ws.max_row, 4)  # 2 en-tête + 2 unités

    def test_round_trip_import(self):
        from collecte.excelio import parse_and_import
        from collecte.models import DataValue
        buf = self._fill_and_get_buffer({
            ("CHU", "SUSP"): 500, ("CHU", "TDRP"): 300,
            ("AGOE", "SUSP"): 520, ("AGOE", "TDRP"): 342,
        })
        report = parse_and_import(buf, self.ds, "202608", save=True)
        self.assertTrue(report.ok)
        self.assertEqual(report.saved, 4)
        self.assertEqual(DataValue.objects.count(), 4)
        dv = DataValue.objects.get(org_unit=self.fs2, data_element=self.de2, period="202608")
        self.assertEqual(dv.value, Decimal("342"))

    def test_import_is_idempotent(self):
        from collecte.excelio import parse_and_import
        from collecte.models import DataValue
        buf1 = self._fill_and_get_buffer({("CHU", "SUSP"): 500})
        parse_and_import(buf1, self.ds, "202608", save=True)
        buf2 = self._fill_and_get_buffer({("CHU", "SUSP"): 600})
        parse_and_import(buf2, self.ds, "202608", save=True)
        qs = DataValue.objects.filter(org_unit=self.fs1, data_element=self.de1, period="202608")
        self.assertEqual(qs.count(), 1)
        self.assertEqual(qs.first().value, Decimal("600"))

    def test_import_detects_non_numeric(self):
        from collecte.excelio import parse_and_import
        buf = self._fill_and_get_buffer({("CHU", "SUSP"): "abc"})
        report = parse_and_import(buf, self.ds, "202608", save=True)
        self.assertFalse(report.ok)
        self.assertEqual(report.skipped, 1)
        self.assertTrue(any("non numérique" in e for e in report.errors))

    def test_import_detects_unknown_org_unit(self):
        from collecte.excelio import build_template, KEY_SEP
        from collecte.excelio import parse_and_import
        from openpyxl import load_workbook
        from io import BytesIO
        data = build_template(self.ds, period="202608")
        wb = load_workbook(BytesIO(data)); ws = wb["Données"]
        # remplacer un code d'unité par un inexistant
        ws.cell(3, 1, "INEXISTANT")
        buf = BytesIO(); wb.save(buf); buf.seek(0)
        report = parse_and_import(buf, self.ds, "202608", save=True)
        self.assertIn("INEXISTANT", report.unknown_org_units)

    def test_dry_run_writes_nothing(self):
        from collecte.excelio import parse_and_import
        from collecte.models import DataValue
        buf = self._fill_and_get_buffer({("CHU", "SUSP"): 500})
        report = parse_and_import(buf, self.ds, "202608", save=False)
        self.assertEqual(report.saved, 1)          # compté
        self.assertEqual(DataValue.objects.count(), 0)  # mais rien écrit

    def test_period_type_mismatch(self):
        from collecte.excelio import parse_and_import
        buf = self._fill_and_get_buffer({("CHU", "SUSP"): 500})
        report = parse_and_import(buf, self.ds, "2026Q3", save=True)  # trimestriel sur formulaire mensuel
        self.assertFalse(report.ok)
        self.assertTrue(any("Période" in e for e in report.errors))

    def test_download_template_http(self):
        url = reverse("collecte:download_template")
        resp = self.client.get(url, {"dataset": self.ds.id})
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])
        self.assertIn("attachment", resp["Content-Disposition"])

    def test_import_page_upload_http(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        buf = self._fill_and_get_buffer({("CHU", "SUSP"): 500, ("AGOE", "SUSP"): 520})
        upload = SimpleUploadedFile(
            "data.xlsx", buf.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        url = reverse("collecte:import_excel")
        resp = self.client.post(url, {
            "dataset": self.ds.id, "period": "202608", "import": "1", "file": upload,
        })
        self.assertEqual(resp.status_code, 200)
        self.assertIsNotNone(resp.context["report"])
        self.assertEqual(resp.context["report"]["saved"], 2)


class FlexibleImportTests(TestCase):
    def setUp(self):
        self.pays = OrgUnit.objects.create(name="Togo", code="TG", level=1)
        self.fs = OrgUnit.objects.create(name="CHU", code="CHU", parent=self.pays)
        self.combo = CategoryCombo.objects.create(name="(défaut)", code="default", is_default=True)
        self.coc = CategoryOptionCombo.objects.create(combo=self.combo, name="(défaut)")
        self.de = DataElement.objects.create(name="Cas suspects", code="SUSP", category_combo=self.combo)
        self.user = User.objects.create_superuser("adm4", "a@b.c", "pwd12345")
        self.client = Client()
        self.client.force_login(self.user)

    def _make_file_with_extra_column(self):
        """Construit un .xlsx avec une colonne connue (SUSP) et une inconnue (NOUVEAU)."""
        from openpyxl import Workbook
        from io import BytesIO
        from collecte.excelio import (KEY_SEP, HEADER_ROW, KEY_ROW, FIRST_DATA_ROW,
                                       OU_CODE_COL, OU_NAME_COL, FIRST_VALUE_COL)
        wb = Workbook(); ws = wb.active; ws.title = "Données"
        ws.cell(HEADER_ROW, OU_CODE_COL, "Code unité")
        ws.cell(HEADER_ROW, OU_NAME_COL, "Unité")
        ws.cell(KEY_ROW, OU_CODE_COL, "__ou_code__")
        ws.cell(KEY_ROW, OU_NAME_COL, "__ou_name__")
        # colonne connue
        ws.cell(HEADER_ROW, FIRST_VALUE_COL, "Cas suspects")
        ws.cell(KEY_ROW, FIRST_VALUE_COL, f"SUSP{KEY_SEP}{self.coc.id}")
        # colonne inconnue
        ws.cell(HEADER_ROW, FIRST_VALUE_COL + 1, "Nouveaux cas graves")
        ws.cell(KEY_ROW, FIRST_VALUE_COL + 1, f"CAS_GRAVES{KEY_SEP}0")
        # une ligne de données
        ws.cell(FIRST_DATA_ROW, OU_CODE_COL, "CHU")
        ws.cell(FIRST_DATA_ROW, OU_NAME_COL, "CHU")
        ws.cell(FIRST_DATA_ROW, FIRST_VALUE_COL, 500)
        ws.cell(FIRST_DATA_ROW, FIRST_VALUE_COL + 1, 12)
        buf = BytesIO(); wb.save(buf); buf.seek(0)
        return buf

    def test_analyze_detects_known_and_unknown(self):
        from collecte.excelio import analyze_columns
        buf = self._make_file_with_extra_column()
        res = analyze_columns(buf, None)
        self.assertIsNone(res["error"])
        self.assertEqual(len(res["known"]), 1)
        self.assertEqual(len(res["unknown"]), 1)
        self.assertEqual(res["unknown"][0]["de_code"], "CAS_GRAVES")

    def test_create_element_from_column(self):
        from collecte.excelio import create_element_from_column
        de = create_element_from_column("Cas graves", "CAS_GRAVES")
        self.assertEqual(de.code, "CAS_GRAVES")
        self.assertTrue(de.category_combo.is_default)

    def test_flexible_flow_create_and_import(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from collecte.models import DataValue

        # Étape 1 : analyse
        buf = self._make_file_with_extra_column()
        upload = SimpleUploadedFile("f.xlsx", buf.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        url = reverse("collecte:import_flexible")
        r1 = self.client.post(url, {"analyze": "1", "period": "202608", "file": upload})
        self.assertEqual(r1.status_code, 200)
        self.assertEqual(r1.context["step"], "review")
        token = r1.context["token"]
        unknown_idx = r1.context["analysis"]["unknown"][0]["index"]

        # Étape 2 : confirmer en créant l'élément inconnu
        r2 = self.client.post(url, {
            "confirm": "1", "token": token, "period": "202608",
            f"action_{unknown_idx}": "create",
            f"name_{unknown_idx}": "Cas graves",
            f"code_{unknown_idx}": "CAS_GRAVES",
        })
        self.assertEqual(r2.status_code, 200)
        self.assertEqual(r2.context["step"], "done")
        # L'élément a été créé
        self.assertTrue(DataElement.objects.filter(code="CAS_GRAVES").exists())
        # Les valeurs (connue + nouvelle) sont importées
        self.assertTrue(DataValue.objects.filter(data_element__code="SUSP", period="202608").exists())
        self.assertTrue(DataValue.objects.filter(data_element__code="CAS_GRAVES", period="202608").exists())

    def test_flexible_flow_ignore_column(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from collecte.models import DataValue

        buf = self._make_file_with_extra_column()
        upload = SimpleUploadedFile("f.xlsx", buf.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        url = reverse("collecte:import_flexible")
        r1 = self.client.post(url, {"analyze": "1", "period": "202608", "file": upload})
        token = r1.context["token"]
        unknown_idx = r1.context["analysis"]["unknown"][0]["index"]
        r2 = self.client.post(url, {
            "confirm": "1", "token": token, "period": "202608",
            f"action_{unknown_idx}": "ignore",
        })
        self.assertEqual(r2.context["step"], "done")
        # L'élément inconnu n'a PAS été créé
        self.assertFalse(DataElement.objects.filter(code="CAS_GRAVES").exists())
        # Mais la colonne connue est bien importée
        self.assertTrue(DataValue.objects.filter(data_element__code="SUSP", period="202608").exists())
