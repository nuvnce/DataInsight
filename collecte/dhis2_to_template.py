"""
dhis2_to_template.py — Convertisseur DHIS2 → modèle de collecte
================================================================

Transforme un export DHIS2 « large » (une ligne = période × formation
sanitaire ; ~270 colonnes désagrégées) en un classeur au **format du modèle de
collecte** (voir `make_template`) : onglets Public / Privé, une ligne par
district × période, colonnes = élément × tranche d'âge.

C'est l'étape qui relie l'export brut au pipeline de fiabilisation. Elle ne
touche pas la base : elle produit un .xlsx intermédiaire, que l'humain relit,
puis importe comme lot (couche 1).

Principe d'agrégation
---------------------
1. Chaque colonne est passée au parseur → (élément, milieu, secteur, âge).
2. On agrège au **district** (orgunitlevel3) × période.
3. Toutes les colonnes qui tombent sur le même (élément, âge) **dans le même
   secteur** sont **sommées**. C'est ainsi que :
     - les colonnes HD_/HH_ « Paludisme grave … » se somment en **HOSP** ;
     - les colonnes H_Info « Décédés (toutes causes) » se somment en **DECES**.
4. Le secteur (Public / Privé) choisit l'onglet de sortie.

Les colonnes non reconnues, ambiguës ou sans dimension sont **signalées** dans
le rapport, pas silencieusement ignorées.
"""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from . import column_parser as CP


PERIOD_COL = "periodcode"
DISTRICT_COL = "orgunitlevel3"
META_COLS = {
    "periodid", "periodname", "periodcode", "perioddescription",
    "orgunitlevel1", "orgunitlevel2", "orgunitlevel3", "orgunitlevel4",
    "orgunitlevel5", "organisationunitid", "organisationunitname",
    "organisationunitcode", "organisationunitdescription",
}

# Éléments et âges du modèle (mêmes codes que make_template / le parseur).
ELEMENTS = ["SUS", "SUS_TESTE", "TDR_FAIT", "TDR_POS", "CONF_TOT",
            "CONF_SIMPLE", "CONF_GRAVE", "GRAVE_REF", "TRAITE_CTA", "HOSP", "DECES"]
AGES = ["TOT", "U5", "O5", "FE"]
SECTORS = [("Public", "PUB"), ("Privé", "PRV")]

NAVY = "1F3864"


@dataclass
class ConvertReport:
    rows_in: int = 0
    districts: int = 0
    periods: int = 0
    columns_ok: int = 0
    columns_unknown: list = field(default_factory=list)
    columns_ambiguous: list = field(default_factory=list)
    columns_missing_dim: list = field(default_factory=list)
    cells_out: int = 0
    errors: list = field(default_factory=list)

    def as_dict(self):
        return self.__dict__.copy()


def convert(path_or_buffer, out_path, sheet_name=0) -> ConvertReport:
    """
    Lit l'export DHIS2, agrège au district, écrit le classeur modèle dans
    `out_path`. Retourne un rapport.
    """
    report = ConvertReport()

    df = pd.read_excel(path_or_buffer, sheet_name=sheet_name)
    df.columns = [str(c).strip() for c in df.columns]
    report.rows_in = len(df)

    if PERIOD_COL not in df.columns or DISTRICT_COL not in df.columns:
        report.errors.append(
            f"Colonnes obligatoires absentes ({PERIOD_COL}, {DISTRICT_COL}). "
            f"Vérifie que l'export contient bien le niveau district."
        )
        return report

    # 1) Plan des colonnes : nom -> (code élément, secteur, âge)
    plan = {}
    for col in df.columns:
        if col in META_COLS:
            continue
        p = CP.parse_column(col)
        if p.status == "unknown_indicator":
            report.columns_unknown.append(col); continue
        if p.status == "ambiguous":
            report.columns_ambiguous.append(col); continue
        if not p.secteur:
            report.columns_missing_dim.append(col); continue
        if p.indicator_code not in ELEMENTS:
            # élément reconnu mais hors modèle (ex. TPI, MID) : on l'ignore
            report.columns_unknown.append(col); continue
        plan[col] = (p.indicator_code, p.secteur, p.tranche)
        report.columns_ok += 1

    if not plan:
        report.errors.append("Aucune colonne exploitable reconnue dans l'export.")
        return report

    # 2) Agrégation : (secteur, district, période, code, âge) -> somme
    #    On somme automatiquement toutes les colonnes qui retombent sur la même
    #    clé (c'est ce qui fusionne HD_/HH_ en HOSP, H_Info en DECES).
    agg = defaultdict(float)
    periods_col = df[PERIOD_COL].astype(str).str.strip()
    district_col = df[DISTRICT_COL].astype(str).str.strip()

    for col, (code, sector, age) in plan.items():
        series = pd.to_numeric(df[col], errors="coerce").fillna(0.0)
        for district, period, val in zip(district_col, periods_col, series):
            if val == 0:
                continue
            agg[(sector, district, period, code, age)] += float(val)

    districts = sorted({k[1] for k in agg})
    periods = sorted({k[2] for k in agg})
    report.districts = len(districts)
    report.periods = len(periods)
    report.cells_out = len(agg)

    # 3) Écriture du classeur au format modèle
    _write_workbook(out_path, agg, districts, periods)
    return report


def _write_workbook(out_path, agg, districts, periods):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for sector_label, sector_code in SECTORS:
        ws = wb.create_sheet(sector_label)
        _write_header(ws)

        r = 3
        for district in districts:
            for period in periods:
                # une ligne seulement si au moins une valeur existe
                row_vals = {}
                has = False
                for code in ELEMENTS:
                    for age in AGES:
                        v = agg.get((sector_code, district, period, code, age))
                        if v:
                            row_vals[(code, age)] = v
                            has = True
                if not has:
                    continue
                ws.cell(row=r, column=1, value=district)
                ws.cell(row=r, column=2, value=_slug(district))
                pcell = ws.cell(row=r, column=3, value=str(period))
                pcell.number_format = "@"
                col = 4
                for code in ELEMENTS:
                    for age in AGES:
                        v = row_vals.get((code, age))
                        if v is not None:
                            # entier si entier, sinon arrondi 2 décimales
                            ws.cell(row=r, column=col,
                                    value=int(v) if float(v).is_integer() else round(v, 2))
                        col += 1
                r += 1

    out = Path(out_path)
    wb.save(out)


def _write_header(ws):
    def style(cell, bg=NAVY, bold=True):
        cell.font = Font(name="Arial", size=10, bold=bold, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, name in enumerate(["District", "Code district", "Période"], start=1):
        style(ws.cell(row=1, column=i, value=name))
        ws.merge_cells(start_row=1, start_column=i, end_row=2, end_column=i)

    col = 4
    for code in ELEMENTS:
        start = col
        for age in AGES:
            style(ws.cell(row=1, column=col, value=code))
            style(ws.cell(row=2, column=col, value=age), bg="2E5496", bold=False)
            col += 1
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=col - 1)
    ws.freeze_panes = "D3"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 12


def _slug(name: str) -> str:
    """Code district : majuscules, sans accents ni espaces superflus."""
    import unicodedata
    n = unicodedata.normalize("NFD", name)
    n = "".join(ch for ch in n if unicodedata.category(ch) != "Mn")
    return n.strip().upper().replace("  ", " ")
